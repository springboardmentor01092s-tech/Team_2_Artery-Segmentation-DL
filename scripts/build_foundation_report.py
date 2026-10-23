"""Build the monochrome foundation milestone report without running ML jobs."""
from pathlib import Path
from PIL import Image as PILImage, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'project_report'
ASSETS = OUT / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
FONT='/System/Library/Fonts/Supplemental/Arial.ttf'
def canvas(height):
    im=PILImage.new('RGB',(1800,height),'white');return im,ImageDraw.Draw(im)
def txt(d,xy,text,size=25):
    d.multiline_text(xy,text,font=ImageFont.truetype(FONT,size),fill='black',anchor='mm',align='center',spacing=10)
def box(d,rect,text,size=25):
    d.rectangle(rect,outline='black',width=3);txt(d,((rect[0]+rect[2])/2,(rect[1]+rect[3])/2),text,size)
def arrow(d,a,b):
    d.line([a,b],fill='black',width=3)
    x,y=b;dx=b[0]-a[0];dy=b[1]-a[1]
    if abs(dx)>=abs(dy):d.polygon([(x,y),(x-12,y-7),(x-12,y+7)],fill='black')
    else:d.polygon([(x,y),(x-7,y-12),(x+7,y-12)],fill='black')

def flow(name, labels):
    im,d=canvas(260);n=len(labels);w=330
    for i, label in enumerate(labels):
        x=10+i*360;box(d,(x,35,x+w,215),label,27)
        if i<n-1:arrow(d,(x+w+3,125),(x+355,125))
    im.save(ASSETS/f'{name}.png')

flow('workflow',['Define scope\nand modality','Acquire and\naudit data','Prepare and\ninspect inputs','Manage cases\nand access','Review research\noutputs'])
flow('preprocessing',['Raw image\nvalidation','Grayscale and\nnormalization','Optional noise\nand contrast','Optional ROI\ncrop','Resize and pad\nmodel input'])

im,d=canvas(620)
for rect,label in [((10,120,430,300),'React frontend\nRole specific views'),((620,120,1090,300),'FastAPI backend\nAuthentication and policy'),((1320,20,1790,200),'SQLAlchemy database\nUsers cases and history'),((1320,350,1790,530),'Protected file storage\nOriginals and previews'),((620,420,1090,600),'Processing services\n2D and separate 3D')]:box(d,rect,label,28)
for a,b in [((430,210),(620,210)),((1090,170),(1320,110)),((1090,250),(1320,440)),((850,300),(850,420)),((1090,510),(1320,480))]:arrow(d,a,b)
im.save(ASSETS/'architecture.png')

im,d=canvas(530)
for offset,vals,title in [(0,[1000,200,300],'Images per task'),(920,[4976,1168,1672],'Syntax polygon annotations')]:
    txt(d,(offset+430,35),title,31);d.line([(offset+65,440),(offset+820,440)],fill='black',width=3)
    txt(d,(offset+40,440),'0',22)
    for i,(label,v) in enumerate(zip(['Train','Validation','Test'],vals)):
        x=offset+125+i*240;y=440-int(v/max(vals)*310)
        d.rectangle((x,y,x+130,440),outline='black',width=3)
        for yy in range(y+10,440,15):d.line([(x+3,yy),(x+127,yy)],fill='black',width=1)
        txt(d,(x+65,y-25),f'{v:,}',27);txt(d,(x+65,480),label,26)
im.save(ASSETS/'eda.png')

im,d=canvas(620)
for rect,label in [((10,10,330,610),'Navigation\n\nDashboard\nCases\nInspection\nSettings'),((360,10,1790,170),'Case reference and permitted actions'),((360,200,1040,610),'Upload and profile\n\nFile selection\nStandard or enhancement\nROI and frame selection\nValidate and upload'),((1070,200,1790,610),'Preview and processing record\n\nOriginal and prepared view\nDimensions and spacing\nWarnings and provenance')]:box(d,rect,label,28)
im.save(ASSETS/'wireframe.png')

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','List Bullet','List Number']:
    s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(4)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Caption'].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('CoroVision AI   |   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(t):doc.add_paragraph(t)
def h(t):doc.add_heading(t,2)
def page(t):doc.add_page_break();doc.add_heading(t,1)
def pic(name,caption,width=6.8):doc.add_picture(str(ASSETS/f'{name}.png'),width=Inches(width));doc.add_paragraph(caption,'Caption')
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for col,width in zip(t.columns,widths):col.width=Inches(width)
    for c,text in zip(t.rows[0].cells,headers):c.text=text
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            tcPr.append(b)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.0
                for r in par.runs:r.font.size=Pt(10);r.bold=ri==0
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.add_heading('CoroVision AI Project Foundation Report',0)
doc.add_paragraph('Artery Segmentation and Stenosis Detection Using Advanced Deep Learning Techniques','Subtitle')
p('Foundation milestone assessment and implementation workflow\n15 October 2026')
h('Purpose and principal finding')
p('CoroVision AI is a research application for organizing de-identified vascular images, preparing reproducible model inputs, and reviewing artery segmentation and model-flagged stenosis regions. The foundation includes a React interface, a FastAPI service, account authentication, authorized case management, image upload, and basic preprocessing. These components support a research demonstration; they do not establish clinical safety or diagnostic accuracy.')
p('The primary imaging workflow uses two-dimensional X-ray coronary angiography. Three-dimensional CT and CTA preparation is a separate research path because image geometry, annotation meaning, and model inputs differ. Dataset preparation is not uniformly complete across every proposed source. ARCADE has detailed EDA and defined preprocessing, while other datasets retain access, licensing, mask interpretation, or patient-grouping conditions.')
h('Project objectives')
p('The first objective is to convert an uploaded study into a traceable prepared image without losing the original. The second is to connect each image to an authorized case, its processing settings, and any recorded review. The third is to maintain separate, reproducible experiments for artery segmentation and stenosis localization instead of treating these as interchangeable labels.')
p('The foundation milestone also defines the application architecture, relational data model, interface structure, and acceptance checks. Success means that an authorized user can sign in, create or access an allowed case, upload a supported image, select a preparation profile, and inspect the resulting preview and processing record. A viewer sees only shared case details, available findings, and assigned clinicians.')
h('Scope of this report')
p('Sections cover objectives and workflow, modality selection, dataset acquisition and annotation structure, exploratory analysis, system architecture, database design, wireframes, application setup, authentication, case management, upload, preprocessing, and expected outcomes. Training completion and claims of highest evaluation metrics are outside this foundation milestone. Training remains paused at the user’s request.')

page('1 Project workflow and imaging modality')
pic('workflow','Figure 1  Foundation workflow with data review before application processing')
h('Define the workflow before implementation')
p('Requirements begin with the intended user and task. A researcher needs reproducible image preparation and access to experiment records. A technician needs intake and quality inspection. A clinical reviewer needs case context and attributed review notes. A standard viewer needs shared case information and results, without model-management controls. These differences determine both navigation and server authorization.')
p('The data workflow starts with an approved source and a preserved raw copy. An audit checks image readability, annotations, split assignments, duplicates, and geometry. Preprocessing then creates a derived input and records its settings. In the application, the case is the organizing unit: the uploaded source, prepared preview, analysis history, and clinical review refer back to that case. A failed validation returns an error rather than silently producing a substitute image.')
h('Selected primary modality')
p('Two-dimensional X-ray coronary angiography is the primary modality because ARCADE provides coronary anatomy and stenosis-region annotations suitable for separate research tasks. Each image is a projection, not a complete three-dimensional artery. Overlapping vessels and acquisition differences limit what can be inferred from one frame. A highlighted region must therefore be described as a model output requiring review, not an established diagnosis or percentage stenosis.')
h('Separate volumetric extension')
p('CT and CTA volumes represent anatomy across spatial slices. Preparing them requires voxel spacing, orientation, origin, and consistent image-mask geometry. Hepatic vessels, aorta, cerebral vessels, and coronary arteries have different labels and scientific questions. They remain separate datasets and model pipelines; their raw samples are not concatenated with coronary X-ray frames into one training tensor.')
p('The application can prepare supported research volumes separately, but a prepared CT volume is not automatically passed to the X-ray inference model. This boundary prevents a technically valid upload from being mistaken for a scientifically valid model input.')

page('2 Dataset selection and annotation structure')
h('Selection and acquisition procedure')
p('A suitable dataset must match the modality, anatomy, and target task. Public availability alone does not confirm permitted reuse or clinically appropriate labels. The acquisition record should identify the source, release version, checksum, licence, annotation definition, and original split. Raw archives and annotations remain unchanged; prepared outputs use separate directories and manifests.')
table(['Dataset or group','Project role and preparation condition'],[
('ARCADE','Primary X-ray benchmark. Syntax and stenosis tasks are separate. Detailed inventory and EDA are available; patient independence remains unresolved.'),
('CADICA','Supplementary coronary lesion research. Local intake records 6,126 selected frames from 42 patients; these are intake counts, not a final independent training cohort.'),
('ICA and DCA1','Supplementary vessel segmentation. DCA1 intake contains 134 pairs. Reuse terms and patient grouping need resolution; ICA intermediate mask semantics also need verification.'),
('Mendeley stenosis','Bounding-box localization resource. CC BY 4.0 was confirmed in the project record; patient grouping remains unresolved. Boxes are not pixel masks.'),
('MSD Task 08','Separate CT hepatic-vessel task. Voxel audit passed for 303 labelled cases; labels distinguish background, vessel, and tumour.'),
('Other CT and brain sources','ImageCAS, ASOCA, AortaSeg-60, TopCoW, and TopBrain remain separate extensions. Acquisition and release-specific terms must be confirmed individually.')],[1.5,5.3])
h('Understand the source structure')
p('ARCADE uses task folders for syntax and stenosis, with train, validation, and test subdivisions. Annotation records refer to an image identifier and category, with polygons and bounding boxes. Polygon coordinates alternate x and y values, while boxes use x, y, width, and height. File names can repeat across tasks and splits, so a safe identifier includes the dataset, task, split, and source image ID.')
p('Rasterized masks are derived training artefacts, not replacements for the original polygons. A binary artery mask discards anatomical class and instance distinctions and must be labelled as a separate target. Likewise, a stenosis box cannot be converted into an expert pixel contour by filling the rectangle. In CT datasets, a tumour class must not be merged into a vessel class without an explicitly justified experiment.')

page('3 Exploratory data analysis')
p('Exploratory data analysis examines the dataset before choosing model assumptions. The ARCADE audit counted images and annotation records, decoded image sizes and modes, checked references, measured intensities and annotation areas, and compared SHA-256 hashes across splits. The following numbers are recorded findings from the existing project EDA, not a fresh scan or new experiment.')
pic('eda','Figure 2  ARCADE split sizes and syntax annotation counts from the recorded EDA')
table(['Task','Train images and polygons','Validation images and polygons','Test images and polygons'],[
('Syntax','1,000 / 4,976','200 / 1,168','300 / 1,672'),('Stenosis','1,000 / 1,625','200 / 406','300 / 386')],[1.0,1.9,2.0,1.9])
h('Interpret the findings')
p('The inventory contains 3,000 image entries across the two tasks and 10,233 polygon annotations. These are not 3,000 unique patients. All listed images are 512 by 512 pixels, but source modes include grayscale and RGB. Consistent grayscale conversion therefore removes a format inconsistency without implying that the underlying acquisition conditions are identical.')
p('Class imbalance is substantial. Syntax category 12 has only one training annotation and is absent from test. Aggregate pixel accuracy can hide failure on small vessels or rare categories. Three unannotated stenosis training frames are retained rather than silently removed; absence of a supplied annotation should not be interpreted as a clinically verified normal study.')
h('Leakage and distribution limitations')
p('The audit found one byte-identical stenosis image shared between train ID 218 and validation ID 155. The two task test sets share 123 exact images. Official splits remain preserved and these findings must accompany evaluation. Hash checks do not detect every related video frame and cannot establish patient-level independence. Split intensity and lesion-size differences also mean a metric from one partition should not be generalized to another without evidence.')

page('4 System architecture')
pic('architecture','Figure 3  Logical architecture with backend controlled access to records and files',6.1)
h('Frontend responsibilities')
p('The React and TypeScript frontend presents role-specific pages and collects user inputs. It displays accessible cases, upload controls, preparation options, previews, and review information. A shared API helper sends requests to the backend with the authenticated bearer token. Interface visibility improves usability, but hiding a button is not treated as a security boundary.')
h('Backend and processing responsibilities')
p('FastAPI receives requests, validates structured inputs, checks the current user and effective permissions, and enforces case-level access. SQLAlchemy manages persistent records. Processing services decode supported images, construct prepared previews, and store provenance. Protected endpoints serve source images, previews, and overlays only after authorization. Images are not exposed simply because their storage paths are known.')
p('The local database defaults to SQLite, which avoids requiring a separate database server during development. PostgreSQL is a deployment target, not evidence that a production database has already been deployed. File storage holds original sources and derived artefacts, while the database keeps their associations and processing metadata. A production migration would require tested backups and coordinated recovery of both database and files.')
h('Research model boundary')
p('Research inference is a downstream service, distinct from validation and preprocessing. Its output should include model version and analysis history, with human review recorded separately. Geometry-aware volume preparation remains independent of the 2D X-ray model. The design does not assume that all acquired vascular datasets support the same inference task.')
h('Deployment considerations')
p('The current development convention uses frontend port 5173 and backend port 8000. Deployment should replace local assumptions with configured API origins, HTTPS, restricted cross-origin settings, secure secrets, and a tested storage policy. Distributed job queues, multi-institution tenancy, and clinical deployment approval are not established by this foundation architecture.')

page('5 Database schema and relationships')
p('The database separates identity, case access, imaging artefacts, and review records. This prevents account information from being mixed into image pixels and supports an audit trail when a case is shared or a result is reviewed. Identifiers link records; filesystem paths point to stored artefacts rather than embedding full images in account tables.')
table(['Entity group','Important fields and purpose'],[
('Users and profiles','User ID, unique email, password hash, role, status, display name, and optional username. Profile data does not independently grant permissions.'),
('Sessions and access policy','Session expiry and revocation; login throttling; role access requests; explicit operational grants with administrator and reason.'),
('Cases and case access','Owner, case code, pseudonymous patient code, modality, anatomy, workflow state, and user-level read or operational access.'),
('Images','Case reference, original and preview paths, source dimensions, available spacing, frame selection, ROI, and preprocessing JSON.'),
('Analysis records','Image reference, model version, mask counts, output paths, validation metadata, and append-only analysis event history.'),
('Clinical workflow','Clinical reviews, report revisions and decisions, case comments, notifications, and audit events.')],[1.65,5.15])
h('Core relationship diagram')
for line in ['User  1 ---- many  Case  1 ---- many  Image', 'User  1 ---- many  CaseAccess  many ---- 1  Case', 'Case  1 ---- many  ClinicalReview and ReportRevision', 'Image  1 ---- many  AnalysisEvent', 'User  1 ---- many  AuthSession and AuditEvent']:
    par=doc.add_paragraph();r=par.add_run(line);r.font.name='Consolas';r.font.size=Pt(10)
p('Ownership and assignment answer different questions. Ownership identifies the creator; a case-access record explicitly shares an existing case with another user. A viewer assignment must not imply upload or analysis permission. The backend combines resource access with the user’s permitted action before allowing a mutation.')
h('Integrity and retention requirements')
p('Foreign keys describe intended relationships, while unique email and username fields prevent account collisions. Duplicate active case codes are checked within an owner’s cases. Source hash plus preparation settings distinguish repeated uploads from intentional variants, but the existing pre-check is not a concurrency-safe unique database constraint. Soft deletion preserves recoverability; retention periods and institution-specific deletion policy require a separate operational decision.')

page('6 Interface wireframes and user experience')
pic('wireframe','Figure 4  Conceptual intake wireframe for a user with upload permission')
h('Authentication and navigation')
p('The sign-in page needs visible identity and password fields, clear validation messages, and an unambiguous sign-in action. Public registration creates a standard viewer account. Professional access is requested and reviewed separately. Institutional SSO is an integration option, but it must not appear to be a functioning university or hospital identity service before a real provider is configured.')
p('Navigation follows responsibilities. A technician focuses on case intake and image inspection; a researcher can inspect preparation records and permitted research results. Clinical reviewers see assigned studies and review actions. A viewer sees a simplified dashboard, assigned cases, and profile settings rather than model-validation statistics or experiment controls. Administrative account management does not automatically grant unrestricted access to case contents.')
h('Intake and quality inspection')
p('An upload screen first establishes the selected case and modality. File selection is followed by a preparation profile and any explicit frame, page, or ROI controls. The interface should show progress during validation and explain an error next to the relevant action. After success, the prepared preview and provenance are available together so an operator can check that the expected source and settings were used.')
p('Standard, Noise Reduction, Contrast Enhancement, and Noise + Contrast are peer choices in one preprocessing profile section. They are not four nested Standard pipelines. Research controls are reserved for authorized roles. Descriptions explain the operation rather than promising that an enhanced-looking image will produce a better model result.')
h('Appearance and accessibility')
p('Appearance preferences are accessed from the bottom of Settings rather than a permanent comparison strip on the main page. Typography, controls, labels, and icons should follow the selected theme consistently. Associated radio labels, keyboard focus, readable contrast, responsive layout, and adequate gaps between sections remain necessary regardless of theme. Wireframes describe structure; they are not a substitute for testing the rendered interface.')

page('7 Frontend backend setup and authentication')
h('Application setup')
p('The frontend uses React, TypeScript, and Vite. FastAPI provides the Python service and generated OpenAPI documentation. SQLAlchemy handles database sessions, and the backend environment contains image-processing and volume-processing dependencies. Local setup should keep dependencies pinned or recorded, use a dedicated environment, and obtain secrets through configuration rather than placing them in client code.')
p('The request helper centralizes backend URLs and bearer-token headers. Development calls use the backend on port 8000. JSON endpoints handle registration, login, case metadata, and reviews; multipart requests carry uploaded files and preparation options. This separation matters because an authentication-only Swagger service cannot replace the case and image-processing API without implementing matching contracts.')
h('Authentication sequence')
p('A user submits an email or username and password. The backend resolves the account, verifies the stored password hash, checks account status and failed-login throttling, and creates a session. A bearer token represents that session in subsequent requests. The current-user endpoint returns account state and effective permissions so the frontend can display appropriate controls.')
p('A protected request still checks the account and session on the server. Token possession does not bypass a suspended account, a revoked session, or a missing case assignment. Password change and sign-out operations revoke applicable sessions. Audit events record account activity. These controls establish a useful prototype baseline but do not constitute a full production security assessment.')
h('Registration and role approval')
p('Public registration grants viewer access rather than accepting an administrator or clinician role selected by the registrant. A user can request an eligible workspace, state the purpose, and wait for administrator review. Professional eligibility is manually verified; the system does not establish medical qualifications from an email address. Sensitive credentials or patient details should not be entered into the request text.')
h('Institutional SSO status')
p('Real single sign-on delegates authentication to a trusted institution identity provider. It requires provider configuration, verified redirect and issuer settings, and a role-mapping policy. A demo username is only a local sample account. It is not real institutional SSO, and connecting an actual hospital or university remains a pending integration.')

page('8 Case management and image upload')
h('Case creation and assignment')
p('An authorized operator creates a case with a unique reference within their active cases, modality, anatomy, and de-identified context. The patient code is pseudonymous and should not be a name or hospital identifier exposed without authorization. Cases keep separate workflow, processing, and report states so a prepared image is not confused with a completed review or approved report.')
p('Case access is explicitly assigned. Reading a shared case does not automatically allow changing metadata, uploading images, running analysis, or exporting results. Clinical reviews carry a reviewer identity, and report approval is a distinct permission. The viewer-facing experience shows available findings and assigned clinicians; it does not display a fabricated doctor or invent a result when the case is still empty.')
h('Upload request and validation')
p('The operator selects a supported raster image, a supported 2D DICOM image or explicit frame, or a PDF page. The backend checks file size, decoding, dimensions, input options, and case authorization. The uploader currently describes a 50 MB limit per input. A file extension alone is not proof of valid content. Unsupported or malformed inputs should be rejected with an actionable explanation.')
p('Original files are retained while derived previews are stored separately. The processing record links source hashes, dimensions, available spacing, selected settings, operations, and warnings. Source-and-settings duplicate checks prevent an identical repeated variant within a case while allowing the same source to be compared under different preparation profiles. Such comparisons are exploratory and must not be used for repeated test-driven model selection.')
h('Protected retrieval and privacy')
p('Preview and source retrieval use protected API endpoints, not public storage links. DICOM identifying-tag checks help detect unsuitable uploads, but automated metadata inspection is not complete de-identification. Operators remain responsible for removing identifying metadata and burned-in text before upload. PDF documents need special attention because visible text can contain identifiers even if image metadata appears clean.')
h('Separate volume upload')
p('The research volume form accepts supported scalar NIfTI inputs and optional aligned masks or a reference volume. It returns a preparation ZIP rather than adding the volume to the 2D image collection. Browser DICOM-series ZIP intake is not implemented; the separate offline importer handles supported regular series. Volume preparation does not start training.')

page('9 Basic preprocessing pipeline')
pic('preprocessing','Figure 5  Actual 2D sequence with optional enhancement steps')
h('Validation and normalization')
p('Validation establishes that the selected image can be decoded and used with the requested options. Conversion produces a consistent grayscale representation. DICOM handling applies modality rescale or lookup information before intensity normalization, while high-bit-depth raster images need an explicit mapping into the preview range. Intensity normalization changes numerical values; it does not recover missing anatomy or correct an incorrect label.')
h('Profile operations')
table(['Profile','Operation before optional ROI and resizing'],[
('Standard','Conservative validation, grayscale conversion, and applicable intensity normalization.'),
('Noise Reduction','Standard preparation plus a 3 by 3 median filter.'),
('Contrast Enhancement','Standard preparation plus clipped autocontrast with a 1 percent cutoff.'),
('Noise + Contrast','Median noise reduction followed by clipped autocontrast.')],[1.75,5.05])
p('The UI path follows the selected profile, and the backend resolves the selected preset into actual operations. Enhancement is optional because median filtering can remove fine structures and contrast changes can alter the model input distribution. The current model’s own input transformation still applies during inference, so these options are not validated accuracy improvements.')
h('ROI resizing and provenance')
p('Region-of-interest extraction uses an explicit user-selected rectangle rather than an invented anatomical detector. The image is resized with aspect ratio preserved and padded onto a 512 by 512 canvas. This avoids stretching anatomy into a square. Original dimensions, crop, resizing, padding, and any available physical spacing are recorded. Missing pixel spacing stays unknown; image dimensions alone cannot supply millimetres.')
p('Quality inspection compares the prepared view against the source. A visually unchanged result can be legitimate for a constant or already high-contrast image. An acceptable output must be readable, correctly oriented for the selected source, associated with the right case, and accompanied by settings and warnings. A successful upload alone is not evidence of clinical suitability.')

page('10 Volume preparation and expected outcomes')
h('Additional processing requirements')
p('Three-dimensional preparation preserves source geometry and checks aligned mask geometry before transforming either. Explicit target spacing controls resampling; categorical masks use nearest-neighbour interpolation to avoid creating fractional class labels. Cropping updates physical origin. Prepared volume arrays and selected slice arrays are separate outputs, with geometry recorded so spatial measurements remain interpretable.')
p('Slice selection chooses a defined frame or spatial plane rather than silently selecting the most attractive result. Paired training augmentation must apply consistent transforms to image and label, use a recorded seed, and remain restricted to training data. It is not automatically applied to an uploaded evaluation case. Registration is used only when a relevant reference is supplied; rigid alignment and optimizer success still require visual quality review.')
h('Expected outcome assessment')
table(['Expected outcome','Current foundation assessment'],[
('Project architecture completed','Core architecture, component responsibilities, and data relationships are defined for the research prototype. Production deployment design remains separate.'),
('Dataset prepared','ARCADE inventory, EDA, and preparation are documented. Broader dataset readiness is conditional on licences, grouping, annotations, and task-specific checks.'),
('Frontend and backend initialized','React and FastAPI implementations are present. A read-only check during this review returned healthy API status and 67 documented paths.'),
('Authentication operational','Local sign-in, sessions, role restrictions, and access requests are implemented. Institutional SSO remains dependent on a real provider.'),
('Image upload operational','Authorized 2D intake, source retention, prepared previews, and processing records are implemented. Complete de-identification is not guaranteed.'),
('Basic preprocessing completed','Standard and optional enhancement profiles are connected to actual operations. Separate geometry-aware volume preparation is available with stated format limits.')],[1.65,5.15])
h('Verification and remaining acceptance checks')
p('The latest recorded targeted regression run passed 29 API, preprocessing, and workflow tests, including pixel-level comparisons of the four presets on a synthetic image. The frontend build passed. These are engineering checks, not a clinical validation study or a new full-system test performed for this report. Browser end-to-end verification of the newly added 3D form remains outstanding.')
p('Before declaring the milestone fully accepted, demonstrate sign-in, allowed and denied case access, valid and invalid upload, every profile, source preservation, and geometry inspection on the intended demonstration environment. Resolve dataset-specific gates before new controlled training. Training remains paused, and neither highest metrics nor whole-body clinical coverage is claimed.')

page('11 Evidence register and handover')
h('Implementation evidence')
p('The report reflects the project workspace as reviewed on 15 October 2026. Current API code takes precedence where older architecture notes describe superseded behaviour. In particular, public registration now creates a viewer account, and an administrator does not automatically receive unrestricted case-content access. No additional dataset download, model training, account registration, or clinical evaluation was performed to prepare this document.')
table(['Evidence location','Use in this report'],[
('backend/app/main.py and frontend/src/App.tsx','API routes, authentication contract, case workflow, image intake, role checks, and current profile choices.'),
('backend/app/models.py and database.py','Entity relationships, persistent records, and local SQLite default.'),
('docs/DATASET_EDA.md','Recorded ARCADE inventory, intensity observations, class imbalance, and duplicate findings.'),
('reports/local_training_intake.json and msd_volume_audit.json','Supplementary intake counts and corrected MSD voxel audit. Older registry acquisition labels are not treated as current.'),
('docs/PREPROCESSING_ENGINE_REQUIREMENTS.md and STANDARD_PREPROCESSING_PIPELINES.md','2D and 3D preparation operations, geometry, limitations, and previous regression evidence.'),
('docs/ROLE_WORKSPACES_IMPLEMENTATION_2026-10-15.md and USER_PAUSED_TRAINING.md','Role-specific interface boundaries, unimplemented integrations, and paused-training constraint.')],[3.4,3.4])
h('Handover artefacts')
p('The foundation handover comprises the application source, backend environment specification, database schema, dataset inventory and audit records, preprocessing implementation, automated tests, and this report. Raw datasets, derived arrays, and checkpoints should be transferred with checksums and access restrictions when moving to a larger computer. Secrets and session tokens should not be bundled into a public project archive.')
h('Recommended next milestone')
p('The next authorized work should focus on an end-to-end demonstration and closure of dataset permission and grouping questions. Once the user explicitly resumes compute work, eligible dataset-specific training can continue from verified checkpoints on suitable hardware. Model selection should use training and validation data, while final test data remains reserved for the declared evaluation. Research outputs continue to require qualified review before any clinical interpretation.')

doc.core_properties.title='CoroVision AI Project Foundation Report'
doc.core_properties.subject='Objectives architecture dataset preparation authentication upload and preprocessing'
doc.core_properties.author='CoroVision AI Project'
path=OUT/'CoroVision_AI_Project_Foundation_Report.docx';doc.save(path)
print(path)
